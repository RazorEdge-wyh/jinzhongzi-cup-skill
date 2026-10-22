#!/usr/bin/env python3
"""Submission linter for Chinese innovation & entrepreneurship competition materials.

Checks (mechanically, not editorially):
  - required sections per material kind (plan / paper / deck)
  - evidence markers: [实测] / [估算] / [待补] balance (plan & deck)
  - process-annotation leakage (e.g. [核实 Sx], 【待填】, "> 自检清单", emoji)
  - label-style headings that make AI text obvious ("定位宣言：" etc.)
  - empty marketing phrases with no evidence ("国际领先", "填补空白", ...)
  - numbers without a nearby source marker
  - citation integrity: in-text [n] vs reference list, type markers, ghost refs
  - figure/table numbering continuity
  - five-dimension coverage keyword scan (plan & deck)
  - judging-dimension mapping table presence (plan)

Usage:
  python check_submission.py <file> [--kind plan|paper|deck] [--json] [--strict]
  python check_submission.py --mapping [--json]     # print a mapping-table skeleton
  python check_submission.py --version

Exit codes: 0 = pass (warnings allowed), 1 = errors (or warnings with --strict).

Stdlib only. Reads .md / .txt / .docx (docx needs python-docx).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

__version__ = "1.2.0"

for _stream in (sys.stdout, sys.stderr):  # Windows console default is GBK
    try:
        _stream.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except Exception:  # pragma: no cover
        pass

# ---------------------------------------------------------------- rules

REQUIRED_SECTIONS = {
    "plan": [
        ("痛点", ("痛点", "需求", "问题")),
        ("证据链", ("证据", "文献", "事实")),
        ("方案", ("方案", "解决")),
        ("差异点", ("不同点", "创新点", "差异")),
        ("市场/产业", ("市场", "产业", "客户")),
        ("团队", ("团队", "组织")),
        ("财务/融资", ("财务", "融资", "成本", "盈利")),
        ("社会价值", ("社会价值", "持续性", "社会责任")),
        ("落地", ("落地", "推广", "湖南")),
    ],
    "paper": [
        ("摘要", ("摘要",)),
        ("引言", ("引言", "前言")),
        ("方法", ("材料与方法", "方法")),
        ("结果", ("结果",)),
        ("讨论", ("讨论",)),
        ("结论", ("结论",)),
        ("参考文献", ("参考文献",)),
    ],
    "deck": [
        ("封面/一句话", ("一句话", "价值主张", "封面")),
        ("痛点", ("痛点", "问题")),
        ("证据", ("证据", "数据")),
        ("方案", ("方案", "产品")),
        ("市场", ("市场", "客户")),
        ("团队", ("团队",)),
        ("财务", ("财务", "成本", "盈利")),
        ("落地", ("落地", "规划")),
    ],
}

FILLER = [
    "国际领先", "国内首创", "填补空白", "颠覆", "革命性", "世界第一", "最先进",
    "巨大市场", "蓝海", "前景广阔", "意义重大", "毫无疑问", "毋庸置疑",
    "显著提升", "大幅提高", "极大改善", "全面提升", "有效解决",
    "众所周知", "不言而喻", "必将",
]

DIMENSIONS = {
    "产品创新与技术突破(25)": ("技术", "创新", "专利", "软著", "壁垒", "迭代", "参数", "样机"),
    "产业认知与市场价值(20)": ("市场", "产业", "客户", "商业模式", "定价", "竞品", "规模", "渠道"),
    "组织架构与团队能力(20)": ("团队", "分工", "导师", "顾问", "股权", "激励"),
    "财务分析与融资规划(20)": ("财务", "收入", "成本", "利润", "现金流", "融资", "资金用途", "预测"),
    "社会价值与持续发展(15)": ("社会价值", "减施", "增收", "死苗", "就业", "生态", "可持续", "湖南"),
}

CITE_INLINE = re.compile(r"\[(\d{1,3})(?:[-,–]\s*\d{1,3})*\]")
REF_ENTRY = re.compile(r"^\s*\[(\d{1,3})\]\s*(.+)$")
TYPE_MARKER = re.compile(r"\[(J|M|D|S|P|C|R|N|EB/OL|DB/OL|Z)\]")
SENTENCE_SPLIT = re.compile(r"(?<=[。；;！!？?])\s*|\n+")
CLAIM_HINT = re.compile(
    r"(\d+(?:\.\d+)?\s*(?:%|％|倍|元|万元|亿元|亩|点|天|CFU|个/g|kg|吨|株|户|台|套|项))"
)
SOURCE_HINT = re.compile(
    r"(\[\d|（来源|\(来源|来源[:：]|文献|据[^，。]{0,12}(?:研究|报道|统计)|数据来源|参见|转引|待实测|待核|估算|假设|目标值)"
)
# left-over process annotations that must not survive into a submission
ANNOTATION_LEAK = re.compile(r"(\[(?:实测|估算|待补|核实[^\]]*|未核实|推断)\]|【(?:待填|待核实|待补)[^】]*】|>\s*自检清单|⚠️|✅|❌|★)")
# "label-style" pseudo-headings that make AI-written text obvious
LABEL_STYLE = re.compile(
    r"^\s*(?:\*\*)?(定位宣言|核心逻辑|一句话概括|系统架构一句话|核心卖点|价值主张一句话|一句话总结)(?:\*\*)?\s*[：:]",
    re.MULTILINE,
)
MAPPING_HINT = re.compile(r"(评分(要点|维度|标准)|得分点).{0,40}(章节|对应|证据)")


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt", ".markdown", ""):
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".docx":
        try:
            from docx import Document  # type: ignore

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        except Exception as exc:  # pragma: no cover
            raise SystemExit(f"无法读取 DOCX（需要 python-docx）：{exc}")
    raise SystemExit(f"不支持的文件类型：{suffix}（支持 .md/.txt/.docx）")


def headings(text: str) -> list[str]:
    out = []
    for line in text.splitlines():
        m = re.match(r"^\s{0,3}(#{1,4})\s+(.*)$", line)
        if m:
            out.append(m.group(2).strip())
    return out


def check_sections(text: str, kind: str) -> tuple[list[str], list[str]]:
    head_text = "\n".join(headings(text)) or text[:4000]
    errors, found = [], []
    for name, keys in REQUIRED_SECTIONS[kind]:
        if any(k in head_text for k in keys):
            found.append(name)
        else:
            errors.append(f"缺少章节：{name}（可接受标题关键词：{'/'.join(keys)}）")
    return errors, found


def check_filler(text: str) -> list[dict]:
    hits = []
    for word in FILLER:
        for m in re.finditer(re.escape(word), text):
            start = max(0, m.start() - 30)
            ctx = text[start : m.end() + 30].replace("\n", " ")
            hits.append({"word": word, "context": ctx.strip()})
    return hits


def check_claims(text: str) -> list[str]:
    missing = []
    for raw in SENTENCE_SPLIT.split(text):
        sentence = raw.strip()
        if len(sentence) < 8 or not CLAIM_HINT.search(sentence):
            continue
        if SOURCE_HINT.search(sentence):
            continue
        if re.search(r"第\s*\d+\s*天|0\s*[—\-–~]\s*30\s*天|v\d|版本", sentence):
            continue
        missing.append(sentence[:120])
    return missing


def check_citations(text: str) -> tuple[list[str], list[str]]:
    errors, warnings = [], []
    lines = text.splitlines()
    ref_idx = None
    for i, line in enumerate(lines):
        if re.match(r"^\s{0,3}#{1,4}\s*参考文献", line) or line.strip() in ("参考文献", "References"):
            ref_idx = i
            break
    refs: dict[int, str] = {}
    if ref_idx is not None:
        for line in lines[ref_idx + 1 :]:
            m = REF_ENTRY.match(line)
            if m:
                refs[int(m.group(1))] = m.group(2)
    # 正文引用只在参考文献之前统计，避免把文末条目自身当成引用
    body_text = "\n".join(lines[:ref_idx]) if ref_idx is not None else text
    inline = {int(m.group(1)) for m in CITE_INLINE.finditer(body_text)}
    if ref_idx is None:
        if inline:
            warnings.append(f"正文有 {len(inline)} 处 [n] 引用，但未找到“参考文献”章节")
        return errors, warnings
    if not refs:
        errors.append("有“参考文献”章节，但没有解析到 [n] 条目")
    missing = sorted(inline - set(refs))
    unused = sorted(set(refs) - inline)
    if missing:
        errors.append(f"正文引用了但文末缺失的序号：{missing}")
    if unused:
        warnings.append(f"文末列出但正文未引用的序号：{unused}（幽灵引用）")
    for num, body in refs.items():
        if not TYPE_MARKER.search(body):
            warnings.append(f"参考文献 [{num}] 缺少文献类型标识（[J]/[M]/[D]/[S]/[P]/[EB/OL] 等）")
    return errors, warnings


def check_figures(text: str) -> list[str]:
    warnings = []
    for label, pattern in (("图", r"图\s*(\d+)"), ("表", r"表\s*(\d+)")):
        nums = sorted({int(x) for x in re.findall(pattern, text)})
        if nums and nums != list(range(1, len(nums) + 1)):
            warnings.append(f"{label}编号不连续或不从 1 开始：{nums}")
    return warnings


def check_dimensions(text: str) -> dict[str, int]:
    return {dim: sum(1 for kw in kws if kw in text) for dim, kws in DIMENSIONS.items()}


def check_length(text: str) -> list[str]:
    warnings = []
    chars = len(re.sub(r"\s", "", text))
    if chars < 1500:
        warnings.append(f"全文仅 {chars} 字（含表格），可能不足以支撑评审")
    long_paras = [p.strip() for p in text.split("\n\n") if len(p) > 600 and not p.strip().startswith("#")]
    if long_paras:
        warnings.append(f"有 {len(long_paras)} 个段落超过 600 字，建议拆分")
    return warnings


def check_markers(text: str) -> tuple[list[str], dict[str, int]]:
    warnings = []
    counts = {k: len(re.findall(rf"\[{k}\]", text)) for k in ("实测", "估算", "待补")}
    if counts["实测"] == 0 and CLAIM_HINT.search(text):
        warnings.append("未发现 [实测] 标记：正文数字可能缺少实测支撑")
    if counts["待补"] > 0:
        warnings.append(f"有 {counts['待补']} 处 [待补]：正式提交前需补齐或改为目标值表述")
    return warnings, counts


def check_annotations(text: str) -> list[dict]:
    hits = []
    for m in ANNOTATION_LEAK.finditer(text):
        start = max(0, m.start() - 25)
        ctx = text[start : m.end() + 25].replace("\n", " ")
        hits.append({"mark": m.group(0), "context": ctx.strip()})
    return hits


def check_label_style(text: str) -> list[str]:
    return [m.group(1) for m in LABEL_STYLE.finditer(text)]


def check_mapping(text: str) -> list[str]:
    if MAPPING_HINT.search(text):
        return []
    return ["未发现“评审维度对照表”（评分要点 ↔ 章节 ↔ 证据）；BP 建议附上（模板见 assets/templates/judging-mapping.md）"]


def mapping_skeleton(found_sections: list[str]) -> str:
    rows = [
        "| 评分维度 | 分值 | 对应章节 | 证据（数字/来源） | 状态 |",
        "|---|---|---|---|---|",
        "| 产品创新与技术突破 | 25 | | | 待填 |",
        "| 产业认知与市场价值 | 20 | | | 待填 |",
        "| 组织架构与团队能力 | 20 | | | 待填 |",
        "| 财务分析与融资规划 | 20 | | | 待填 |",
        "| 社会价值与持续发展 | 15 | | | 待填 |",
    ]
    return "\n".join(rows) + "\n\n已检测到的章节：" + ("、".join(found_sections) or "（无）")


# ---------------------------------------------------------------- main

def main() -> int:
    ap = argparse.ArgumentParser(description="Competition material submission linter")
    ap.add_argument("file", nargs="?", type=Path)
    ap.add_argument("--kind", choices=sorted(REQUIRED_SECTIONS), default="plan")
    ap.add_argument("--json", action="store_true", help="输出 JSON")
    ap.add_argument("--strict", action="store_true", help="有警告即返回 1")
    ap.add_argument("--mapping", action="store_true", help="输出评审维度对照表骨架")
    ap.add_argument("--version", action="version", version=f"check_submission.py {__version__}")
    args = ap.parse_args()

    if args.mapping:
        print(mapping_skeleton([]))
        return 0
    if args.file is None:
        ap.error("需要 <file> 或 --mapping")
    if not args.file.is_file():
        raise SystemExit(f"文件不存在：{args.file}")

    text = read_text(args.file)
    sec_errors, sec_found = check_sections(text, args.kind)
    filler = check_filler(text)
    claims = check_claims(text)
    cite_errors, cite_warnings = check_citations(text)
    annotations = check_annotations(text)
    labels = check_label_style(text)
    marker_warnings, marker_counts = check_markers(text)
    dims = check_dimensions(text) if args.kind != "paper" else {}

    errors = list(sec_errors) + list(cite_errors)
    warnings = (
        [f"空话/无证据最高级：{w}" for w in {h["word"] for h in filler}]
        + [f"数字无来源：{c}" for c in claims[:12]]
        + [f"正文批注残留（提交前清理）：{a['mark']}" for a in annotations[:8]]
        + [f"标签体句式（AI 感明显）：{lb}：" for lb in labels]
        + marker_warnings
        + cite_warnings
        + check_figures(text)
        + check_length(text)
        + (check_mapping(text) if args.kind == "plan" else [])
        + [f"五维覆盖偏弱：{d}（命中 {n} 个关键词）" for d, n in dims.items() if n <= 2]
    )

    if marker_counts:
        warnings.append(
            f"证据标记统计：[实测] {marker_counts['实测']} / [估算] {marker_counts['估算']} / [待补] {marker_counts['待补']}"
        )

    report = {
        "file": str(args.file),
        "kind": args.kind,
        "chars": len(re.sub(r"\s", "", text)),
        "sections_found": sec_found,
        "dimension_coverage": dims,
        "markers": marker_counts,
        "errors": errors,
        "warnings": warnings,
        "filler_count": len(filler),
        "unsourced_claims": len(claims),
        "annotation_leaks": len(annotations),
        "label_style_hits": len(labels),
    }

    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"== 自检：{args.file}（类型 {args.kind}，{report['chars']} 字）")
        print("\n[章节]")
        print("  已找到：" + ("、".join(sec_found) or "（无）"))
        if dims:
            print("\n[五维覆盖（命中关键词数）]")
            for d, n in dims.items():
                print(f"  {d:<26} {n}")
        if errors:
            print("\n[错误]")
            for e in errors:
                print("  ✗ " + e)
        if warnings:
            print("\n[警告]")
            for w in warnings:
                print("  ! " + w)
        if not errors and not warnings:
            print("\n通过：未发现问题。")
        elif not errors:
            print(f"\n通过（{len(warnings)} 条警告，建议逐条处理）。")
        else:
            print(f"\n未通过：{len(errors)} 个错误，{len(warnings)} 条警告。")

    return 1 if errors or (args.strict and warnings) else 0


if __name__ == "__main__":
    sys.exit(main())
